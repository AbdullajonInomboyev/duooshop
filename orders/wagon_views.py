"""
Yuk xati (накладная) — Chop etish (HTML/Print) va Excel (.xlsx) eksport tizimi.
1. waybill_single — Bitta buyurtma yuk xati (HTML yoki Excel).
2. waybill_grouped — Kunlik jamlangan yuk xati (Jami yuklash va Do'konlar taqsimoti).
3. waybill_multiple — Bir nechta buyurtmani bitta varaqda chop etish/yuklash.
4. waybill_by_brand — Ishlab chiqaruvchi/brend bo'yicha saralangan yuk xati.
5. waybill_selected_by_brand — Tanlangan do'konlar tovarlarini brendlarga bo'lib berish.
"""
from collections import defaultdict, OrderedDict
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.contrib.admin.views.decorators import staff_member_required
from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse
from django.utils import timezone

from catalog.models import Brand
from .models import Order, OrderItem


def _get_excel_border():
    return Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )


# ============================================================
# 1. BITTA BUYURTMA YUK XATI
# ============================================================
@staff_member_required
def waybill_single(request, order_id):
    """Bitta buyurtma yuk xati. Guruhlangan tovarlar bilan (HTML yoki ?export=excel)."""
    order = get_object_or_404(
        Order.objects.select_related("shop", "user"), pk=order_id
    )
    items = order.items.select_related(
        "product", "product__category", "product__brand"
    ).all()

    # Bir xil tovarlarni (nom + brend) guruhlash
    grouped = {}
    for it in items:
        cat = it.product.category.name if it.product and it.product.category else ""
        brand = it.product.brand.name if it.product and it.product.brand else ""
        key = (it.product_name, brand)
        if key in grouped:
            grouped[key]["qty"] += it.quantity
            grouped[key]["total"] += it.line_total
        else:
            grouped[key] = {
                "category": cat,
                "brand": brand,
                "name": it.product_name,
                "unit": getattr(it, "unit", "") or getattr(it.product, "unit", "dona"),
                "price": it.unit_price,
                "qty": it.quantity,
                "total": it.line_total,
            }

    rows = []
    grand_total = 0
    for i, data in enumerate(
        sorted(grouped.values(), key=lambda x: (x["category"], x["brand"], x["name"])), 1
    ):
        rows.append({"n": i, **data})
        grand_total += data["total"]

    # EXCEL EKSPORT
    if request.GET.get("export") == "excel":
        wb = openpyxl.Workbook()
        ws = wb.active
        order_num = getattr(order, "receipt_number", None) or str(order.id)
        ws.title = "Yuk xati"
        thin_border = _get_excel_border()
        shop = order.shop

        # Shapka — HTML bilan bir xil: korxona, manzil, telefon, sana, № (7 ustun)
        ws.merge_cells("A1:G1")
        tc = ws["A1"]
        tc.value = f"YUK XATI № {order_num}"
        tc.font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
        tc.fill = PatternFill(start_color="009D4D", fill_type="solid")
        tc.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        district = shop.district.name if shop and shop.district else ""
        addr = (shop.address_text if shop and shop.address_text else "")
        full_addr = district + (", " + addr if addr else "")
        phone = shop.owner.phone if shop and shop.owner else "-"
        ws["A3"] = "Korxona nomi:"; ws["B3"] = shop.name if shop else "-"
        ws["A4"] = "Adres:"; ws["B4"] = full_addr
        ws["A5"] = "Telefoni:"; ws["B5"] = phone
        ws["A6"] = "Sana:"; ws["B6"] = order.created_at.strftime("%d.%m.%Y %H:%M")
        for r in range(3, 7):
            ws[f"A{r}"].font = Font(bold=True)

        # Jadval sarlavhasi — HTML bilan bir xil (Kategoriya + Brend bor)
        headers = ["№", "Kategoriya", "Brend", "Mahsulot nomi",
                   "Narxi", "Miqdori", "Summasi"]
        hdr_row = 8
        ws.row_dimensions[hdr_row].height = 24
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=hdr_row, column=col_idx, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="009D4D", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border

        cur = hdr_row + 1
        for r in rows:
            ws.cell(row=cur, column=1, value=r["n"]).alignment = Alignment(horizontal="center")
            ws.cell(row=cur, column=2, value=r.get("category", "")).font = Font(bold=True)
            ws.cell(row=cur, column=3, value=r.get("brand", ""))
            ws.cell(row=cur, column=4, value=r["name"])
            ws.cell(row=cur, column=5, value=float(r["price"])).number_format = "#,##0"
            ws.cell(row=cur, column=6, value=r["qty"]).alignment = Alignment(horizontal="center")
            ws.cell(row=cur, column=7, value=float(r["total"])).number_format = "#,##0"
            for col_i in range(1, 8):
                ws.cell(row=cur, column=col_i).border = thin_border
            cur += 1

        # JAMI qatori (sariq fon, HTML dagidek)
        ws.merge_cells(start_row=cur, start_column=1, end_row=cur, end_column=6)
        jami_c = ws.cell(row=cur, column=1, value="JAMI")
        jami_c.font = Font(bold=True)
        jami_c.alignment = Alignment(horizontal="right")
        jami_c.fill = PatternFill(start_color="FFF44F", fill_type="solid")
        tot_c = ws.cell(row=cur, column=7, value=f"=SUM(G{hdr_row+1}:G{cur-1})")
        tot_c.font = Font(bold=True)
        tot_c.number_format = "#,##0"
        tot_c.fill = PatternFill(start_color="FFF44F", fill_type="solid")
        for col_i in range(1, 8):
            ws.cell(row=cur, column=col_i).border = thin_border

        # Ustun kengliklari
        widths = [6, 18, 16, 32, 12, 10, 14]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="nakladnoy_{order_num}.xlsx"'
        wb.save(resp)
        return resp

    context = {
        "order": order,
        "shop": order.shop,
        "rows": rows,
        "grand_total": grand_total,
        "date": order.created_at,
    }
    return render(request, "admin/waybill_single.html", context)


# ============================================================
# 2. SANA BO'YICHA JAMLANGAN YUK XATI
# ============================================================
@staff_member_required
def waybill_grouped(request):
    """Sana bo'yicha jamlangan yuk xati (HTML yoki ?export=excel)."""
    date_str = request.GET.get("date")
    if date_str:
        try:
            day = datetime.strptime(date_str, "%Y-%m-%d").date()
        except ValueError:
            day = timezone.now().date()
    else:
        day = timezone.now().date()

    orders = (
        Order.objects.filter(created_at__date=day)
        .exclude(status="cancelled")
        .select_related("shop", "user")
        .prefetch_related("items__product", "items__product__category", "items__product__brand")
    )

    totals = defaultdict(lambda: {"qty": 0, "category": "", "brand": "", "unit": ""})
    by_shop = OrderedDict()

    for order in orders:
        shop_name = order.shop.name if order.shop else str(order.user)
        if shop_name not in by_shop:
            by_shop[shop_name] = []

        for it in order.items.all():
            cat = it.product.category.name if it.product and it.product.category else ""
            brand = it.product.brand.name if it.product and it.product.brand else ""
            unit_val = getattr(it, "unit", "") or getattr(it.product, "unit", "dona")
            key = (it.product_name, brand)

            totals[key]["qty"] += it.quantity
            totals[key]["category"] = cat
            totals[key]["brand"] = brand
            totals[key]["unit"] = unit_val

            by_shop[shop_name].append({
                "category": cat,
                "brand": brand,
                "name": it.product_name,
                "qty": it.quantity,
                "unit": unit_val,
                "total": it.line_total,
            })

    summary = []
    for i, ((name, brand), data) in enumerate(
        sorted(totals.items(), key=lambda x: (x[1]["category"], x[1]["brand"], x[0][0])), 1
    ):
        summary.append({
            "n": i,
            "category": data["category"],
            "brand": data["brand"],
            "name": name,
            "qty": data["qty"],
            "unit": data["unit"],
        })

    # EXCEL EKSPORT
    if request.GET.get("export") == "excel":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Jamlangan_{day.strftime('%d.%m')}"
        thin_border = _get_excel_border()

        ws.merge_cells("A1:E1")
        title_cell = ws["A1"]
        title_cell.value = f"JAMLANGAN TOVARLAR RO'YXATI — {day.strftime('%d.%m.%Y')}"
        title_cell.font = Font(size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="17A2B8", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        headers = ["№", "Kategoriya", "Brend", "Mahsulot nomi", "Jami miqdor"]
        ws.row_dimensions[3].height = 22
        for idx, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=idx, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="343A40", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")

        cur = 4
        for item in summary:
            ws.cell(row=cur, column=1, value=item["n"]).alignment = Alignment(horizontal="center")
            ws.cell(row=cur, column=2, value=item["category"])
            ws.cell(row=cur, column=3, value=item["brand"])
            ws.cell(row=cur, column=4, value=item["name"])
            ws.cell(row=cur, column=5, value=f"{item['qty']} {item['unit']}").alignment = Alignment(horizontal="right")
            for ci in range(1, 6):
                ws.cell(row=cur, column=ci).border = thin_border
            cur += 1

        for col in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 14)

        resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="jamlangan_{day}.xlsx"'
        wb.save(resp)
        return resp

    context = {
        "day": day,
        "summary": summary,
        "by_shop": by_shop,
        "order_count": orders.count(),
        "shop_count": len(by_shop),
    }
    return render(request, "admin/waybill_grouped.html", context)


# ============================================================
# 3. BIR NECHTA BUYURTMANI CHOP ETISH
# ============================================================
@staff_member_required
def waybill_multiple(request):
    """Bir nechta tanlangan buyurtmalar uchun yuk xati (?ids=1,2,3)."""
    ids_str = request.GET.get("ids", "")
    ids = [int(x) for x in ids_str.split(",") if x.strip().isdigit()]
    orders = (
        Order.objects.filter(pk__in=ids)
        .exclude(status="cancelled")
        .select_related("shop", "user")
        .order_by("shop__name")
    )

    waybills = []
    for order in orders:
        items = order.items.select_related("product", "product__category", "product__brand").all()
        grouped = {}
        for it in items:
            cat = it.product.category.name if it.product and it.product.category else ""
            brand = it.product.brand.name if it.product and it.product.brand else ""
            key = (it.product_name, brand)
            if key in grouped:
                grouped[key]["qty"] += it.quantity
                grouped[key]["total"] += it.line_total
            else:
                grouped[key] = {
                    "category": cat,
                    "brand": brand,
                    "name": it.product_name,
                    "price": it.unit_price,
                    "qty": it.quantity,
                    "total": it.line_total,
                }
        rows = []
        grand_total = 0
        for i, data in enumerate(
            sorted(grouped.values(), key=lambda x: (x["category"], x["brand"], x["name"])), 1
        ):
            rows.append({"n": i, **data})
            grand_total += data["total"]

        waybills.append({
            "order": order,
            "shop": order.shop,
            "rows": rows,
            "grand_total": grand_total,
            "date": order.created_at,
        })

    return render(request, "admin/waybill_multiple.html", {"waybills": waybills})


# ============================================================
# 4. BREND BO'YICHA YUK XATI
# ============================================================
@staff_member_required
def waybill_by_brand(request):
    """Ishlab chiqaruvchi/brend bo'yicha jamlangan yuk xati (?date=YYYY-MM-DD&brand=<id>)."""
    date_str = request.GET.get("date")
    if date_str:
        try:
            day = datetime.strptime(date_str, "%Y-%m-%d").date()
        except ValueError:
            day = timezone.now().date()
    else:
        day = timezone.now().date()

    brand_id = request.GET.get("brand")
    brands = Brand.objects.filter(is_active=True).order_by("name")
    selected_brand = None
    summary = []

    if brand_id:
        selected_brand = Brand.objects.filter(pk=brand_id).first()
        items = (
            OrderItem.objects.filter(order__created_at__date=day, product__brand_id=brand_id)
            .exclude(order__status="cancelled")
            .select_related("product", "product__category")
        )
        totals = defaultdict(lambda: {"qty": 0, "category": "", "unit": ""})
        for it in items:
            cat = it.product.category.name if it.product and it.product.category else ""
            unit_val = getattr(it, "unit", "") or getattr(it.product, "unit", "dona")
            totals[it.product_name]["qty"] += it.quantity
            totals[it.product_name]["category"] = cat
            totals[it.product_name]["unit"] = unit_val

        for i, (name, data) in enumerate(
            sorted(totals.items(), key=lambda x: (x[1]["category"], x[0])), 1
        ):
            summary.append({
                "n": i,
                "category": data["category"],
                "name": name,
                "qty": data["qty"],
                "unit": data["unit"],
            })

    context = {
        "day": day,
        "brands": brands,
        "selected_brand": selected_brand,
        "summary": summary,
    }
    return render(request, "admin/waybill_by_brand.html", context)


# ============================================================
# 5. TANLANGAN BUYURTMALARNI BRENDLARGA BO'LISH
# ============================================================
@staff_member_required
def waybill_selected_by_brand(request):
    """Tanlangan buyurtmalardan brendlar bo'yicha guruhlangan yuk xati (?ids=1,2,3)."""
    ids_str = request.GET.get("ids", "")
    ids = [int(x) for x in ids_str.split(",") if x.strip().isdigit()]
    orders = (
        Order.objects.filter(pk__in=ids)
        .exclude(status="cancelled")
        .select_related("shop")
        .prefetch_related("items__product", "items__product__category", "items__product__brand")
    )

    by_brand = OrderedDict()
    for order in orders:
        for it in order.items.all():
            brand = it.product.brand.name if it.product and it.product.brand else "Boshqa"
            cat = it.product.category.name if it.product and it.product.category else ""
            unit_val = getattr(it, "unit", "") or getattr(it.product, "unit", "dona")

            if brand not in by_brand:
                by_brand[brand] = {}

            key = it.product_name
            if key in by_brand[brand]:
                by_brand[brand][key]["qty"] += it.quantity
            else:
                by_brand[brand][key] = {
                    "category": cat,
                    "name": it.product_name,
                    "qty": it.quantity,
                    "unit": unit_val,
                }

    brands = []
    for brand_name in sorted(by_brand.keys()):
        rows = []
        for i, data in enumerate(
            sorted(by_brand[brand_name].values(), key=lambda x: (x["category"], x["name"])), 1
        ):
            rows.append({"n": i, **data})
        brands.append({"brand": brand_name, "rows": rows})

    return render(request, "admin/waybill_selected_by_brand.html", {
        "brands": brands,
        "order_count": orders.count(),
    })